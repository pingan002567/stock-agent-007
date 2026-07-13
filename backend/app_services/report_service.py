from __future__ import annotations

from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any
from uuid import uuid4

from backend.app_services.audit_service import AuditService
from backend.app_services.context_builder import ContextBuilder
from backend.app_services.decision_journal_service import DecisionJournalService
from backend.app_services.monitor_service import MonitorService
from backend.app_services.strategy_service import StrategyService
from backend.persistence.file_store import FileStore
from backend.persistence.repositories import WorkbenchRepository
from backend.schemas import (
    AuthorityLevel,
    BacktestRun,
    EventContext,
    PaperPortfolioProjection,
    PaperPortfolioSnapshot,
    Report,
    ReportGenerateRequest,
    ReportQualityCheck,
    ReportTemplate,
    StockContext,
    model_to_dict,
)
from backend.stock_domain.report_tools import generate_stock_dashboard
from backend.stock_domain import report_charts as rc


def _utc_now() -> datetime:
    return datetime.now(timezone.utc)


REPORT_TEMPLATE_REGISTRY: tuple[ReportTemplate, ...] = (
    ReportTemplate(
        template_id="stock_research_default",
        report_type="stock_research",
        name="个股研究模板",
        summary="基于 StockContext 生成研究结论、证据、行业与竞争格局、反对理由与风险提示。",
        source_types=["stock"],
        sections=["结论", "理由", "行业与竞争格局", "反对理由", "证据引用", "免责声明"],
    ),
    ReportTemplate(
        template_id="monitor_review_default",
        report_type="monitor_review",
        name="盯盘复盘模板",
        summary="基于已有 monitor event 生成事件复盘、证据引用与候选动作。",
        source_types=["monitor_event"],
        sections=["事件概览", "触发说明", "证据引用", "候选动作", "执行约束"],
    ),
    ReportTemplate(
        template_id="strategy_backtest_default",
        report_type="strategy_backtest",
        name="策略回测模板",
        summary="基于已有 BacktestRun 生成回测摘要、风险总结与候选动作。",
        source_types=["backtest_run"],
        sections=["回测概览", "关键指标", "风险总结", "候选动作", "执行约束"],
    ),
    ReportTemplate(
        template_id="paper_portfolio_review_default",
        report_type="paper_portfolio_review",
        name="Paper Portfolio 复盘模板",
        summary="基于已有 Paper Portfolio snapshot 生成调仓效果复盘，不重新拉取 live quote。",
        source_types=["paper_portfolio_snapshot"],
        sections=["组合概览", "持仓变化", "风险策略引用", "Warnings", "执行约束"],
        version="v0.16",
    ),
)


class ReportService:
    def __init__(
        self,
        *,
        repo: WorkbenchRepository,
        context_builder: ContextBuilder,
        monitor_service: MonitorService,
        strategy_service: StrategyService,
        audit_service: AuditService,
        file_store: FileStore,
        decision_journal_service: DecisionJournalService,
    ) -> None:
        self.repo = repo
        self.context_builder = context_builder
        self.monitor_service = monitor_service
        self.strategy_service = strategy_service
        self.audit_service = audit_service
        self.file_store = file_store
        self.decision_journal_service = decision_journal_service
        self.repo.seed_report_templates(REPORT_TEMPLATE_REGISTRY)

    def list_templates(self) -> list[ReportTemplate]:
        persisted = {item.template_id: item for item in self.repo.list_report_templates(visible_only=False)}
        items: list[ReportTemplate] = []
        for template in REPORT_TEMPLATE_REGISTRY:
            stored = persisted.get(template.template_id)
            items.append(template.model_copy(update={"visible": stored.visible if stored else template.visible}))
        return [item for item in items if item.visible]

    def create_stock_report(self, context: StockContext, report_type: str = "深研") -> Report:
        title = f"{context.symbol} {report_type}报告"
        return self.generate(
            ReportGenerateRequest(
                report_type="stock_research",
                source_type="stock",
                source_id=context.symbol,
                title=title,
            )
        )

    def generate(self, request: ReportGenerateRequest) -> Report:
        template = self._resolve_template(request)
        source = self._resolve_source(request)
        report = self._compose_report(request, template, source)
        check = self._build_quality_check(report)
        report = report.model_copy(
            update={
                "quality_status": check.status,
                "quality_summary": check.summary,
                "latest_quality_check_id": check.check_id,
            }
        )
        self._write_markdown(report)
        self.repo.save_report(report)
        self.repo.save_report_quality_check(check)
        self.decision_journal_service.auto_link_report(report)
        self.audit_service.record(
            "report generated",
            f"{report.report_type} source={report.source_type}:{report.source_id} report={report.report_id}",
            AuthorityLevel.A2,
        )
        return report

    def list_reports(self) -> list[Report]:
        return self.repo.list_reports()

    def get_report(self, report_id: str) -> Report:
        report = self.repo.get_report(report_id)
        if not report:
            raise KeyError(report_id)
        return report

    def export_report(self, report_id: str) -> dict[str, Any]:
        report = self.get_report(report_id)
        self.audit_service.record("report export", report_id, AuthorityLevel.A2)
        return {
            "report_id": report.report_id,
            "format": "markdown",
            "content": report.content,
            "quality_status": report.quality_status,
            "quality_summary": report.quality_summary,
        }

    def export_report_audio(self, report_id: str) -> bytes:
        """报告 → 语音（m4a/AAC）。零依赖方案：macOS ``say``（有中文语音）。

        say 不可用（非 macOS / 无中文语音）时 raise RuntimeError，路由层转 501。
        产物同时留档 file_store（reports/<id>.m4a），供重复播放不重复合成。
        """
        import re
        import shutil
        import subprocess
        import tempfile

        report = self.get_report(report_id)

        cached = Path(self.file_store.root) / "reports" / f"{report_id}.m4a"
        if cached.is_file():
            return cached.read_bytes()

        if shutil.which("say") is None:
            raise RuntimeError("audio export requires macOS `say` (not found on this host)")

        voices = subprocess.run(
            ["say", "-v", "?"], capture_output=True, text=True, timeout=15
        ).stdout
        zh_voices = [line.split()[0] for line in voices.splitlines() if "zh_CN" in line]
        preferred = next((v for v in ("Tingting", "Meijia") if v in zh_voices), None)
        voice = preferred or (zh_voices[0] if zh_voices else None)
        if voice is None:
            raise RuntimeError("no zh_CN voice installed for `say`")

        speech = self._markdown_to_speech_text(report.content)
        self.audit_service.record("report audio export", report_id, AuthorityLevel.A2)

        with tempfile.TemporaryDirectory(prefix="report-audio-") as tmp:
            text_path = Path(tmp) / "speech.txt"
            out_path = Path(tmp) / "out.m4a"
            text_path.write_text(speech, encoding="utf-8")
            proc = subprocess.run(
                ["say", "-v", voice, "-o", str(out_path), "--data-format=aac", "-f", str(text_path)],
                capture_output=True, text=True, timeout=300,
            )
            if proc.returncode != 0 or not out_path.is_file():
                raise RuntimeError(f"say failed: {proc.stderr.strip() or proc.returncode}")
            audio = out_path.read_bytes()

        cached.parent.mkdir(parents=True, exist_ok=True)
        cached.write_bytes(audio)
        return audio

    @staticmethod
    def _markdown_to_speech_text(content: str, max_chars: int = 6000) -> str:
        """Markdown → 朗读文本：去表格线/标记符/引用标记，保留标题与句子。"""
        import re

        lines: list[str] = []
        for raw in content.splitlines():
            line = raw.strip()
            if not line or set(line) <= {"-", "|", ":", " ", "="}:
                continue
            if line.startswith("|"):
                cells = [
                    re.sub(r"[▲▼█░●◆*_`]+", "", c).strip()
                    for c in line.strip("|").split("|")
                ]
                cells = [c for c in cells if c and set(c) - {"-", ":", "—"}]
                if cells:
                    lines.append("，".join(cells) + "。")
                continue
            line = re.sub(r"^#+\s*", "", line)          # 标题井号
            line = re.sub(r"[*_`>]+", "", line)          # 强调/引用符
            line = re.sub(r"\[来源:[^\]]*\]", "", line)  # 引用标记不朗读
            line = re.sub(r"[▲▼█░●◆]+", "", line)       # 图形字符
            line = line.strip()
            if line:
                lines.append(line)
        return "\n".join(lines)[:max_chars]

    def export_report_pptx(self, report_id: str) -> bytes:
        """报告 → 路演 PPT（python-pptx，随 markitdown 已在依赖内）。

        结构映射：H1→标题页；每个 H2 节→内容页（md 表格→pptx 表格，列表→bullet）；
        每页最多 12 行，超出自动续页。
        """
        from io import BytesIO

        from pptx import Presentation
        from pptx.util import Inches, Pt

        report = self.get_report(report_id)
        self.audit_service.record("report pptx export", report_id, AuthorityLevel.A2)

        prs = Presentation()
        prs.slide_width = Inches(13.33)
        prs.slide_height = Inches(7.5)

        title_slide = prs.slides.add_slide(prs.slide_layouts[0])
        title_slide.shapes.title.text = report.title
        subtitle = title_slide.placeholders[1]
        subtitle.text = (
            f"{report.symbol or ''} · {report.report_type}"
            f" · {report.created_at[:10]}\n{report.disclaimer or '仅供研究，不构成投资建议。'}"
        )

        def clean(text: str) -> str:
            import re

            text = re.sub(r"[*_`>]+", "", text)
            return re.sub(r"\[来源:[^\]]*\]", "", text).strip()

        # 解析 markdown：按 H2 分节；节内区分表格块与文本行
        sections: list[tuple[str, list[Any]]] = []
        current: tuple[str, list[Any]] | None = None
        table_buf: list[list[str]] = []

        def flush_table() -> None:
            nonlocal table_buf
            if table_buf and current is not None:
                current[1].append(("table", table_buf))
            table_buf = []

        for raw in report.content.splitlines():
            line = raw.strip()
            if line.startswith("## "):
                flush_table()
                current = (clean(line[3:]), [])
                sections.append(current)
                continue
            if current is None or not line or line.startswith("# ") or set(line) <= {"-", "="}:
                continue
            if line.startswith("|"):
                cells = [clean(c) for c in line.strip("|").split("|")]
                if cells and not all(set(c) <= {"-", ":", " "} for c in cells if c):
                    table_buf.append(cells)
                continue
            flush_table()
            text = clean(line.lstrip("-").strip() if line.startswith("- ") else line)
            if text:
                current[1].append(("bullet" if line.startswith("- ") else "text", text))
        flush_table()

        blank_layout = prs.slide_layouts[6]
        for section_title, blocks in sections:
            page_blocks: list[list[Any]] = [[]]
            weight = 0
            for block in blocks:
                rows = len(block[1]) if block[0] == "table" else 1
                if weight + rows > 12 and page_blocks[-1]:
                    page_blocks.append([])
                    weight = 0
                page_blocks[-1].append(block)
                weight += rows
            for page_index, page in enumerate(page_blocks):
                slide = prs.slides.add_slide(blank_layout)
                heading = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12), Inches(0.9))
                run = heading.text_frame.paragraphs[0].add_run()
                run.text = section_title if page_index == 0 else f"{section_title}（续）"
                run.font.size = Pt(28)
                run.font.bold = True

                top = 1.4
                for kind, payload in page:
                    if kind == "table":
                        rows, cols = len(payload), max(len(r) for r in payload)
                        height = min(0.4 * rows, 5.6)
                        shape = slide.shapes.add_table(
                            rows, cols, Inches(0.6), Inches(top), Inches(12), Inches(height)
                        )
                        for r, row in enumerate(payload):
                            for c in range(cols):
                                cell = shape.table.cell(r, c)
                                cell.text = row[c] if c < len(row) else ""
                                cell.text_frame.paragraphs[0].font.size = Pt(13)
                                if r == 0:
                                    cell.text_frame.paragraphs[0].font.bold = True
                        top += height + 0.25
                    else:
                        box = slide.shapes.add_textbox(Inches(0.7), Inches(top), Inches(11.8), Inches(0.45))
                        paragraph = box.text_frame.paragraphs[0]
                        run = paragraph.add_run()
                        run.text = ("• " if kind == "bullet" else "") + payload
                        run.font.size = Pt(15)
                        top += 0.45
                    if top > 6.8:
                        break

        buffer = BytesIO()
        prs.save(buffer)
        return buffer.getvalue()

    def export_report_pdf(self, report_id: str) -> bytes:
        from fpdf import FPDF  # listed in core deps; raises ImportError if absent

        report = self.get_report(report_id)
        self.audit_service.record("report pdf export", report_id, AuthorityLevel.A2)

        font_path = str(Path(__file__).resolve().parent.parent / "assets" / "fonts" / "NotoSansSC-Regular.ttf")
        pdf = FPDF(format="A4")
        pdf.set_auto_page_break(auto=True, margin=18)
        pdf.add_page()
        # Register the same Unicode TTF for every style so any bold/italic
        # request can never fall back to the latin-1-only core font (helvetica).
        for style in ("", "B", "I", "BI"):
            pdf.add_font("NotoSansSC", style, font_path)
        pdf.set_font("NotoSansSC", "", 11)

        def clean(text: str) -> str:
            # Strip inline emphasis/code markers; bold is conveyed via heading size.
            return text.replace("**", "").replace("`", "")

        def write(text: str, size: float, lead: float, *, gap: float = 0.0, style: str = "") -> None:
            pdf.set_font("NotoSansSC", style, size)
            # new_x=LMARGIN so consecutive multi_cells reset to the left margin
            # (default RIGHT leaves x at the page edge → "no horizontal space").
            pdf.multi_cell(0, lead, text, new_x="LMARGIN", new_y="NEXT")
            if gap:
                pdf.ln(gap)

        def is_sep(text: str) -> bool:
            t = text.strip()
            return bool(t) and set(t) <= {"-", "|", ":", " "}

        def render_table(block: list[str]) -> None:
            rows: list[list[str]] = []
            for ln in block:
                if is_sep(ln):
                    continue
                cells = [clean(c.strip()) for c in ln.strip().strip("|").split("|")]
                rows.append(cells)
            if not rows:
                return
            pdf.set_font("NotoSansSC", "", 9.5)
            with pdf.table(
                borders_layout="MINIMAL",
                line_height=pdf.font_size * 1.7,
                first_row_as_headings=True,
                cell_fill_color=245,
                cell_fill_mode="ROWS",
            ) as table:
                for r in rows:
                    trow = table.row()
                    for cell in r:
                        trow.cell(cell)
            pdf.set_font("NotoSansSC", "", 11)
            pdf.ln(2)

        lines = report.content.split("\n")
        i, n = 0, len(lines)
        while i < n:
            stripped = lines[i].strip()
            if not stripped:
                pdf.ln(2)
                i += 1
                continue
            # Table block: a "|" row immediately followed by a separator row.
            if stripped.startswith("|") and i + 1 < n and is_sep(lines[i + 1]) and "|" in lines[i + 1]:
                block: list[str] = []
                while i < n and lines[i].strip().startswith("|"):
                    block.append(lines[i])
                    i += 1
                render_table(block)
                continue
            # Horizontal rule.
            if set(stripped) <= {"-", "="} and len(stripped) >= 3:
                y = pdf.get_y() + 1
                pdf.set_draw_color(200, 200, 200)
                pdf.line(pdf.l_margin, y, pdf.w - pdf.r_margin, y)
                pdf.ln(4)
                i += 1
                continue
            body = clean(stripped)
            if stripped.startswith("> "):
                write(body[2:], 11, 6, style="I", gap=1)
            elif stripped.startswith("### "):
                write(body[4:], 13, 7, gap=1, style="B")
            elif stripped.startswith("## "):
                write(body[3:], 15, 8, gap=1.5, style="B")
            elif stripped.startswith("# "):
                write(body[2:], 18, 10, gap=2, style="B")
            elif stripped.startswith(("- ", "* ", "+ ")):
                write("• " + body[2:], 11, 6)
            elif len(body) >= 2 and body.startswith("_") and body.endswith("_"):
                write(body.strip("_"), 10, 5, style="I")
            else:
                write(body, 11, 6)
            i += 1

        return bytes(pdf.output())

    def get_quality(self, report_id: str) -> dict[str, Any]:
        report = self.get_report(report_id)
        items = self.repo.list_report_quality_checks(report.report_id)
        latest = items[0] if items else None
        return {
            "report_id": report.report_id,
            "quality_status": report.quality_status,
            "latest": model_to_dict(latest) if latest else None,
            "items": [model_to_dict(item) for item in items],
        }

    def rerun_quality(self, report_id: str) -> dict[str, Any]:
        report = self.get_report(report_id)
        check = self._build_quality_check(report)
        self.repo.save_report_quality_check(check)
        updated = report.model_copy(
            update={
                "quality_status": check.status,
                "quality_summary": check.summary,
                "latest_quality_check_id": check.check_id,
            }
        )
        self.repo.save_report(updated)
        self.audit_service.record("report quality rerun", report_id, AuthorityLevel.A2)
        return {
            "report_id": updated.report_id,
            "quality_status": updated.quality_status,
            "latest": model_to_dict(check),
        }

    def _resolve_template(self, request: ReportGenerateRequest) -> ReportTemplate:
        by_id = {item.template_id: item for item in REPORT_TEMPLATE_REGISTRY}
        if request.template_id:
            template = by_id.get(request.template_id)
            if not template:
                raise KeyError(f"unknown report template: {request.template_id}")
            if request.report_type != template.report_type:
                raise ValueError(
                    f"template/report_type mismatch: template={template.template_id} "
                    f"expects {template.report_type}, got {request.report_type}"
                )
        else:
            template = next((item for item in REPORT_TEMPLATE_REGISTRY if item.report_type == request.report_type), None)
            if not template:
                raise KeyError(f"unknown report type: {request.report_type}")
        if request.source_type not in template.source_types:
            raise ValueError(f"template {template.template_id} does not support source_type={request.source_type}")
        return template

    def _resolve_source(self, request: ReportGenerateRequest) -> dict[str, Any]:
        if request.source_type == "stock":
            context = self.context_builder.build_stock_context(request.source_id)
            return {
                "source_id": context.symbol,
                "source_label": f"{context.symbol} {context.name}",
                "symbol": context.symbol,
                "payload": context,
                "evidence_refs": ["stock_context", "provider_router"],
            }
        if request.source_type == "monitor_event":
            event = self._resolve_monitor_event(request.source_id)
            return {
                "source_id": event.event_id,
                "source_label": event.title,
                "symbol": event.symbol,
                "payload": event,
                "evidence_refs": ["monitor_event", "monitor_status"],
            }
        if request.source_type == "backtest_run":
            run = self._resolve_backtest_run(request.source_id)
            return {
                "source_id": run.run_id,
                "source_label": run.strategy_name,
                "symbol": run.universe[0] if run.universe else "",
                "payload": run,
                "evidence_refs": ["backtest_run", "strategy_spec"],
            }
        if request.source_type == "paper_portfolio_snapshot":
            snapshot = self.repo.get_paper_portfolio_snapshot(request.source_id)
            if not snapshot:
                raise KeyError(request.source_id)
            return {
                "source_id": snapshot.snapshot_id,
                "source_label": f"Paper Snapshot {snapshot.snapshot_id}",
                "symbol": "",
                "payload": snapshot,
                "evidence_refs": ["paper_portfolio_snapshot", "paper_portfolio_projection"],
            }
        raise KeyError(f"unsupported source_type: {request.source_type}")

    def _resolve_monitor_event(self, source_id: str) -> EventContext:
        if source_id == "latest":
            items = self.monitor_service.list_events(limit=1)
            if not items:
                raise KeyError("monitor event not found")
            return items[0]
        event = self.repo.get_monitor_event(source_id)
        if event:
            return event
        items = self.monitor_service.list_events(limit=20)
        match = next((item for item in items if item.event_id == source_id), None)
        if not match:
            raise KeyError(source_id)
        return match

    def _resolve_backtest_run(self, source_id: str) -> BacktestRun:
        if source_id == "latest":
            runs = self.repo.list_backtest_runs(limit=1)
            if not runs:
                raise KeyError("backtest run not found")
            return runs[0]
        return self.strategy_service.get_backtest(source_id)

    def _compose_report(self, request: ReportGenerateRequest, template: ReportTemplate, source: dict[str, Any]) -> Report:
        payload = source["payload"]
        if request.report_type == "stock_research":
            return self._compose_stock_research(request, template, source, payload)
        if request.report_type == "monitor_review":
            return self._compose_monitor_review(request, template, source, payload)
        if request.report_type == "strategy_backtest":
            return self._compose_strategy_backtest(request, template, source, payload)
        if request.report_type == "paper_portfolio_review":
            return self._compose_paper_portfolio_review(request, template, source, payload)
        raise KeyError(f"unsupported report_type: {request.report_type}")

    def _compose_stock_research(
        self,
        request: ReportGenerateRequest,
        template: ReportTemplate,
        source: dict[str, Any],
        context: StockContext,
    ) -> Report:
        dashboard = generate_stock_dashboard(context, mode="research")
        report_id = f"report_{context.symbol.lower()}_{uuid4().hex[:8]}"
        title = request.title or f"{context.symbol} 研究报告"

        price = context.price
        holding = context.holding
        ai = context.ai_state
        tech = dashboard["technical_analysis"]
        advice = dashboard["holding_advice"]
        stance = dashboard["stance_summary"]

        content = "\n".join(
            [
                f"# {title}",
                "",
                f"**{context.name}（{context.symbol}）** · {context.market} · {context.sector or '—'}"
                f" / {context.industry or '—'}",
                "",
                f"> **结论**：{dashboard['conclusion']}（置信度 {ai.confidence}，有效期 {stance['valid_for']}）",
                "",
                "## 评级总览",
                *rc.table(
                    ["维度", "评估", "可视化"],
                    [
                        ["AI 评分", f"{ai.score} / 100", rc.bar(ai.score, 100)],
                        ["立场", ai.stance, ""],
                        ["风险标签", ai.risk_label, ""],
                        ["置信度", ai.confidence, ""],
                        ["权限级别", stance["authority_level"], ""],
                    ],
                ),
                "",
                "## 行情快照",
                *rc.table(
                    ["指标", "数值"],
                    [
                        ["最新价", rc.fmt(price.last)],
                        ["涨跌幅", rc.pct_marker(price.change_pct)],
                        ["数据源", price.source + (" · 已降级" if price.degraded else "")],
                        ["更新时间", price.updated_at],
                    ],
                ),
                "",
                "## 风险雷达",
                *rc.table(
                    ["风险维度", "强度", "评分"],
                    [[item["name"], rc.bar(item["value"], 100), item["value"]] for item in dashboard["risk_bars"]],
                ),
                "",
                "## 技术面",
                *rc.table(
                    ["项", "判断"],
                    [
                        ["趋势", tech["trend"]],
                        ["动能", tech["momentum"]],
                        ["支撑位", rc.fmt(tech["support"])],
                        ["压力位", rc.fmt(tech["resistance"])],
                    ],
                ),
                "",
                "## 行业与竞争格局",
                *self._industry_section(context),
                "",
                "## 持仓建议",
                *rc.table(
                    ["项", "数值"],
                    [
                        ["当前仓位", rc.gauge(holding.weight_pct, 100, suffix="%")],
                        ["建议仓位", rc.gauge(advice["suggested_weight"], 100, suffix="%")],
                        ["操作", advice["action"]],
                        ["持仓盈亏", rc.pct_marker(holding.pnl_pct) if holding.pnl_pct is not None else "—"],
                    ],
                ),
                f"\n{advice['reason']}",
                "",
                "## 核心理由",
                *[f"- {item}" for item in dashboard["reasons"]],
                "",
                "## 反对理由 / 风险提示",
                *[f"- {item}" for item in dashboard["counter_reasons"]],
                "",
                "## 证据引用",
                *[f"- `{item}`" for item in source["evidence_refs"]],
                "",
                "---",
                f"_{dashboard['disclaimer']}_",
            ]
        )
        return Report(
            report_id=report_id,
            title=title,
            symbol=context.symbol,
            report_type=request.report_type,
            conclusion=dashboard["conclusion"],
            evidence_count=len(dashboard["reasons"]),
            content=content,
            template_id=template.template_id,
            template_name=template.name,
            source_type=request.source_type,
            source_id=source["source_id"],
            source_label=source["source_label"],
            evidence_refs=list(dict.fromkeys([*source["evidence_refs"], "report_template:stock_research_default"])),
            valid_until=(_utc_now() + timedelta(days=1)).isoformat(),
            disclaimer=dashboard["disclaimer"],
            degraded=context.price.degraded,
            degraded_reason=context.price.degraded_reason,
            payload={"dashboard": dashboard},
        )

    def _compose_monitor_review(
        self,
        request: ReportGenerateRequest,
        template: ReportTemplate,
        source: dict[str, Any],
        event: EventContext,
    ) -> Report:
        explanation = self.monitor_service.explain_event(event=event)
        candidate_actions = [{"action": item, "source": "monitor_event"} for item in event.suggested_actions]
        candidate_action_lines = [f"- {item['action']}" for item in candidate_actions] or ["- 无"]
        report_id = f"report_monitor_{event.symbol.lower()}_{uuid4().hex[:8]}"
        title = request.title or f"{event.symbol} 盯盘复盘报告"
        execution_guard = {
            "research_only": True,
            "auto_trade": False,
            "place_real_order_enabled": False,
        }
        severity_marker = {"high": "● 高", "medium": "● 中", "low": "● 低"}.get(
            str(event.severity).lower(), f"● {event.severity}"
        )
        evidence_rows = [
            [item.get("type") or item.get("ref") or "evidence", rc.fmt(item.get("value")) if item.get("value") is not None else (item.get("detail") or item.get("note") or "—")]
            for item in event.evidence
        ] or [["—", "无证据记录"]]
        content = "\n".join(
            [
                f"# {title}",
                "",
                f"> **{event.title}** · 标的 `{event.symbol}` · 严重度 {severity_marker}",
                "",
                "## 事件概览",
                *rc.table(
                    ["项", "值"],
                    [
                        ["事件 ID", f"`{event.event_id}`"],
                        ["标的", event.symbol],
                        ["严重度", severity_marker],
                        ["触发规则", event.trigger_rule or "—"],
                        ["触发时间", event.triggered_at],
                    ],
                ),
                "",
                "## 触发说明",
                explanation["summary"],
                "",
                "## 证据引用",
                *rc.table(["证据", "明细"], evidence_rows),
                *([f"- `{item}`" for item in source["evidence_refs"]]),
                "",
                "## 候选动作",
                *candidate_action_lines,
                "",
                "## 执行约束",
                "- `execution_guard.auto_trade = false`（研究态，禁止自动下单）",
                "",
                "---",
                "_仅供研究，不构成投资建议。_",
            ]
        )
        return Report(
            report_id=report_id,
            title=title,
            symbol=event.symbol,
            report_type=request.report_type,
            conclusion=explanation["summary"],
            evidence_count=len(event.evidence),
            content=content,
            template_id=template.template_id,
            template_name=template.name,
            source_type=request.source_type,
            source_id=source["source_id"],
            source_label=source["source_label"],
            evidence_refs=list(
                dict.fromkeys(
                    [*source["evidence_refs"], *[item.get("type") or item.get("ref") or "evidence" for item in event.evidence]]
                )
            ),
            valid_until=(_utc_now() + timedelta(days=1)).isoformat(),
            candidate_actions=candidate_actions,
            execution_guard=execution_guard,
            payload={"event": model_to_dict(event), "explanation": explanation},
        )

    def _compose_strategy_backtest(
        self,
        request: ReportGenerateRequest,
        template: ReportTemplate,
        source: dict[str, Any],
        run: BacktestRun,
    ) -> Report:
        candidate_actions = [dict(item) for item in run.candidate_actions]
        execution_guard = dict(run.execution_guard)
        execution_guard["auto_trade"] = False
        candidate_action_lines = [f"- {self._format_candidate_action(item)}" for item in candidate_actions] or ["- 无"]
        sample_size = run.metrics.get("sample_size") or len(run.universe)
        conclusion = (
            f"{run.strategy_name} 回测样本 {sample_size}，"
            f"{'结果降级，需复核数据来源。' if run.degraded else '可继续作为研究输入。'}"
        )
        report_id = f"report_backtest_{uuid4().hex[:8]}"
        title = request.title or f"{run.strategy_name} 回测报告"
        period = run.period or {}
        equity_curve = [
            sig.get("equity") or sig.get("nav") or sig.get("value")
            for sig in run.signals
            if isinstance(sig, dict)
        ]
        spark = rc.sparkline([v for v in equity_curve if v is not None])
        content = "\n".join(
            [
                f"# {title}",
                "",
                f"> **结论**：{conclusion}",
                "",
                "## 回测概览",
                *rc.table(
                    ["项", "值"],
                    [
                        ["策略", f"{run.strategy_name}（{run.strategy_type}）"],
                        ["run_id", f"`{run.run_id}`"],
                        ["标的池", ", ".join(run.universe) if run.universe else "—"],
                        ["区间", f"{period.get('start', '—')} → {period.get('end', '—')}"],
                        ["样本数", sample_size],
                        ["数据状态", "已降级，需复核" if run.degraded else "正常"],
                    ],
                ),
                *([f"\n净值走势：{spark}", ""] if spark else [""]),
                "## 关键指标",
                *rc.table(
                    ["指标", "数值", "可视化"],
                    [[key, rc.fmt(value), self._metric_bar(key, value)] for key, value in run.metrics.items()],
                ),
                "",
                "## 风险总结",
                *rc.table(
                    ["风险项", "数值"],
                    [[key, rc.fmt(value)] for key, value in run.risk_summary.items()],
                ),
                "",
                "## 候选动作",
                *candidate_action_lines,
                "",
                "## 执行约束",
                "- `execution_guard.auto_trade = false`（研究态，禁止自动下单）",
                "",
                "---",
                "_仅供研究，不构成投资建议。_",
            ]
        )
        return Report(
            report_id=report_id,
            title=title,
            symbol=source["symbol"],
            report_type=request.report_type,
            conclusion=conclusion,
            evidence_count=len(run.evidence_refs),
            content=content,
            template_id=template.template_id,
            template_name=template.name,
            source_type=request.source_type,
            source_id=source["source_id"],
            source_label=source["source_label"],
            evidence_refs=list(dict.fromkeys([*source["evidence_refs"], *run.evidence_refs])),
            valid_until=(_utc_now() + timedelta(days=1)).isoformat(),
            disclaimer="仅供研究，不构成投资建议。",
            degraded=run.degraded,
            degraded_reason=run.degraded_reason,
            candidate_actions=candidate_actions,
            execution_guard=execution_guard,
            risk_policy_ref=run.risk_policy_ref,
            payload={
                "strategy_snapshot": run.strategy_snapshot,
                "period": run.period,
                "metrics": run.metrics,
                "signals": run.signals,
                "risk_summary": run.risk_summary,
                "risk_policy_ref": model_to_dict(run.risk_policy_ref) if run.risk_policy_ref else None,
            },
        )

    def _compose_paper_portfolio_review(
        self,
        request: ReportGenerateRequest,
        template: ReportTemplate,
        source: dict[str, Any],
        snapshot: PaperPortfolioSnapshot,
    ) -> Report:
        projection = PaperPortfolioProjection(**snapshot.payload)
        report_id = f"report_paper_portfolio_{uuid4().hex[:8]}"
        title = request.title or "Paper Portfolio 调仓复盘"
        execution_guard = {
            "research_only": True,
            "auto_trade": False,
            "place_real_order_enabled": False,
        }
        warning_lines = [f"- `{item.code}`：{item.message}" for item in projection.warnings] or ["- 无"]
        position_rows = [
            [
                item.symbol,
                rc.fmt(item.quantity),
                rc.fmt(item.market_value),
                rc.pct_marker(item.pnl) if isinstance(item.pnl, (int, float)) and abs(item.pnl) <= 100 else rc.fmt(item.pnl),
                rc.gauge(item.weight_pct, 100, width=10, suffix="%"),
            ]
            for item in projection.positions[:10]
        ] or [["—", "—", "—", "—", "无持仓"]]
        risk_lines = [
            f"- `{item.policy_id}` @v{item.version} · updated_at={item.updated_at}"
            for item in projection.risk_policy_refs
        ] or ["- 无冻结 risk policy ref"]
        invested = (snapshot.market_value or 0)
        content = "\n".join(
            [
                f"# {title}",
                "",
                f"> **净值估算 {rc.fmt(snapshot.equity_estimate)}** · "
                f"浮动盈亏 {rc.pct_marker((snapshot.pnl_estimate / invested * 100) if invested else 0)}"
                f"（{rc.fmt(snapshot.pnl_estimate)}）",
                "",
                "## 组合概览",
                *rc.table(
                    ["项", "值"],
                    [
                        ["快照 ID", f"`{snapshot.snapshot_id}`"],
                        ["基线 ID", f"`{snapshot.baseline_id}`"],
                        ["截至", snapshot.as_of],
                        ["净值估算", rc.fmt(snapshot.equity_estimate)],
                        ["现金估算", rc.fmt(snapshot.cash_estimate)],
                        ["持仓市值", rc.fmt(snapshot.market_value)],
                        ["浮动盈亏", rc.fmt(snapshot.pnl_estimate)],
                        ["数据状态", "已降级，需复核" if snapshot.degraded else "正常"],
                    ],
                ),
                "",
                "## 持仓明细",
                *rc.table(["标的", "数量", "市值", "盈亏", "权重"], position_rows),
                "",
                "## 风险策略引用",
                *risk_lines,
                "",
                "## 预警",
                *warning_lines,
                "",
                "## 执行约束",
                "- `execution_guard.auto_trade = false`（研究态，禁止自动下单）",
                "",
                "---",
                "_仅供研究，不构成投资建议。_",
            ]
        )
        return Report(
            report_id=report_id,
            title=title,
            symbol="",
            report_type=request.report_type,
            conclusion=f"Paper Portfolio equity={snapshot.equity_estimate} pnl={snapshot.pnl_estimate}",
            evidence_count=len(projection.positions) + len(projection.warnings),
            content=content,
            template_id=template.template_id,
            template_name=template.name,
            source_type=request.source_type,
            source_id=source["source_id"],
            source_label=source["source_label"],
            evidence_refs=list(
                dict.fromkeys(
                    [
                        *source["evidence_refs"],
                        *projection.quotes.keys(),
                        *[item.code for item in projection.warnings],
                    ]
                )
            ),
            valid_until=(_utc_now() + timedelta(days=1)).isoformat(),
            disclaimer="仅供研究，不构成投资建议。",
            degraded=snapshot.degraded,
            execution_guard=execution_guard,
            risk_policy_ref=projection.latest_risk_policy_ref,
            payload={
                "snapshot": model_to_dict(snapshot),
                "projection": snapshot.payload,
            },
        )

    def _industry_section(self, context: StockContext) -> list[str]:
        """行业与竞争格局章节：非降级给硬数据表格，降级给明确说明（不编数据）。"""
        from backend.stock_domain.industry_tools import get_industry_context

        try:
            ctx = get_industry_context(symbol=context.symbol)
        except Exception as exc:
            return [f"> 行业数据获取异常：{exc}（本节降级，建议人工核验行业地位）"]
        if ctx.get("degraded"):
            return [f"> 行业数据不可用：{ctx.get('reason') or '未知原因'}（本节降级，壁垒/上下游请结合研报与公开信息人工判断）"]

        snapshot = ctx.get("snapshot") or {}
        valuation = ctx.get("valuation") or {}
        target = ctx.get("target") or {}
        lines: list[str] = [
            f"所属行业：**{ctx['industry']}**（{ctx.get('company_count')} 家公司）"
            f" · 行业涨跌 {rc.pct_marker(snapshot.get('change_pct')) if snapshot.get('change_pct') is not None else '—'}",
            "",
            *rc.table(
                ["行业内指标", "数值"],
                [
                    [
                        "市值排名（推算口径）",
                        f"{target.get('cap_rank')} / {target.get('cap_rank_total')}"
                        if target.get("cap_rank")
                        else "—",
                    ],
                    [
                        "PE 分位 vs 行业",
                        f"{target.get('pe_percentile')}%（行业中位 {valuation.get('pe_median')}）"
                        if target.get("pe_percentile") is not None
                        else "—",
                    ],
                    [
                        "PB 分位 vs 行业",
                        f"{target.get('pb_percentile')}%（行业中位 {valuation.get('pb_median')}）"
                        if target.get("pb_percentile") is not None
                        else "—",
                    ],
                ],
            ),
            "",
            "行业市值 Top5：",
            *[
                f"- {item['name']}（{item['symbol']}）PE {rc.fmt(item.get('pe'))} / 涨跌 {rc.pct_marker(item.get('change_pct')) if item.get('change_pct') is not None else '—'}"
                for item in (ctx.get("top_constituents") or [])[:5]
            ],
            "",
            "> 市值为推算口径（成交额/换手率），仅用于行业内相对排序 `[来源: industry]`；"
            "壁垒/上下游属定性判断，需结合研报与公开信息，本模板不自动生成。",
        ]
        return lines

    def _build_quality_check(self, report: Report) -> ReportQualityCheck:
        issues: list[dict[str, Any]] = []
        if not report.content.strip().startswith("# "):
            issues.append({"level": "error", "code": "missing_heading", "message": "Markdown 缺少一级标题。"})
        if not report.evidence_refs:
            issues.append({"level": "error", "code": "missing_evidence_refs", "message": "报告缺少 evidence_refs。"})
        if "仅供研究" not in report.disclaimer and "仅供研究" not in report.content:
            issues.append({"level": "error", "code": "missing_disclaimer", "message": "报告缺少研究免责声明。"})
        if report.candidate_actions and report.execution_guard.get("auto_trade") is not False:
            issues.append(
                {
                    "level": "error",
                    "code": "candidate_actions_guard",
                    "message": "含 candidate_actions 的报告必须显式声明 auto_trade=false。",
                }
            )
        if report.degraded:
            issues.append(
                {
                    "level": "warning",
                    "code": "degraded_source",
                    "message": report.degraded_reason or "源数据降级，结论需要人工复核。",
                }
            )
        if not report.valid_until:
            issues.append({"level": "warning", "code": "missing_valid_until", "message": "报告缺少有效期。"})
        if report.report_type == "stock_research":
            if "## 行业与竞争格局" not in report.content:
                issues.append(
                    {
                        "level": "warning",
                        "code": "missing_industry_section",
                        "message": "个股研究报告缺少「行业与竞争格局」章节。",
                    }
                )
            elif "行业数据不可用" in report.content or "行业数据获取异常" in report.content:
                issues.append(
                    {
                        "level": "warning",
                        "code": "industry_section_degraded",
                        "message": "行业章节数据降级，行业地位结论需人工复核。",
                    }
                )
        errors = [item for item in issues if item["level"] == "error"]
        warnings = [item for item in issues if item["level"] == "warning"]
        if errors:
            status = "failed"
        elif warnings:
            status = "warning"
        else:
            status = "passed"
        score = max(0, 100 - len(errors) * 35 - len(warnings) * 10)
        summary = f"{status}: {len(errors)} errors, {len(warnings)} warnings"
        return ReportQualityCheck(
            check_id=f"quality_{uuid4().hex[:12]}",
            report_id=report.report_id,
            template_id=report.template_id,
            report_type=report.report_type,
            source_type=report.source_type,
            source_id=report.source_id,
            status=status,
            summary=summary,
            score=score,
            issues=issues,
            evidence_refs=report.evidence_refs,
            degraded=report.degraded,
        )

    def _write_markdown(self, report: Report) -> None:
        path = self.file_store.write_text("reports", f"{report.report_id}.md", report.content)
        report.markdown_path = str(path)

    def _metric_bar(self, key: str, value: Any) -> str:
        """Best-effort 0..100 gauge for a backtest metric; empty when not mappable."""
        if isinstance(value, bool) or not isinstance(value, (int, float)):
            return ""
        import math

        if not math.isfinite(value):
            return ""
        k = key.lower()
        if "sharpe" in k or "sortino" in k:
            return rc.bar(value, vmax=3.0)
        if "drawdown" in k:
            mag = abs(value) * 100 if abs(value) <= 1 else abs(value)
            return rc.bar(mag, vmax=100)
        if -1.0 <= value <= 1.0:
            return rc.bar(value * 100, vmax=100)
        if any(tag in k for tag in ("rate", "win", "pct", "ratio")) and 0 <= value <= 100:
            return rc.bar(value, vmax=100)
        return ""

    def _format_candidate_action(self, action: dict[str, Any]) -> str:
        label = action.get("action") or action.get("type") or "candidate_action"
        symbol = action.get("symbol")
        delta = action.get("delta_weight_pct")
        bits = [str(label)]
        if symbol:
            bits.append(str(symbol))
        if delta is not None:
            bits.append(f"delta={delta}")
        return " | ".join(bits)
